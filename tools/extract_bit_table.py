import openpyxl, json, sys
src, out, units = sys.argv[1], sys.argv[2], sys.argv[3]
wb = openpyxl.load_workbook(src, data_only=True); ws = wb.active
groups = []
for c in range(10, 84, 15):
    m = ws.cell(12, c+1).value; op = ws.cell(12, c+3).value
    if m: groups.append((c, m, op))
keys = ["rpm","feed","stepdown","stepover","ramp_rpm","ramp_feed","plunge_rpm","plunge_feed","leadin_rpm","leadin_feed","leadout_rpm","leadout_feed","retract_rpm","retract_feed"]
def num(v):
    try: return float(v)
    except: return None
tools=[]; skipped=[]
for r in range(14, ws.max_row+1):
    pid = ws.cell(r,1).value; desc = ws.cell(r,2).value
    if not pid or not desc or num(ws.cell(r,4).value) is None:
        if pid or desc: skipped.append((r,pid,desc))
        continue
    link = ws.cell(r,9).hyperlink.target if ws.cell(r,9).hyperlink else None
    t = dict(product_id=str(pid).strip(), description=str(desc).strip(), shank_dia=num(ws.cell(r,3).value),
             cut_dia=num(ws.cell(r,4).value), tip_dia=num(ws.cell(r,5).value), flutes=int(num(ws.cell(r,6).value) or 0),
             flute_len=num(ws.cell(r,7).value), coating=ws.cell(r,8).value, link=link, cutting=[])
    for c,m,op in groups:
        vals = [num(ws.cell(r,c+1+i).value) for i in range(14)]
        if vals[0] and vals[1]:
            d = {k:v for k,v in zip(keys,vals) if v is not None}; d.update(material=m, operation=op); t["cutting"].append(d)
    tools.append(t)
json.dump(dict(source="SainSmart Genmitsu bit table V6.0 Dec 2022 by Graham Bland (community data, used with attribution)", units=units,
  machine=dict(type=ws.cell(11,2).value, max_rpm=ws.cell(11,3).value, min_rpm=ws.cell(11,5).value, max_feed=ws.cell(11,7).value, spindle_w=ws.cell(11,9).value),
  tools=tools), open(out,"w"), indent=1)
print(len(tools), "tools; groups", groups); print("skipped", skipped[:30])
from collections import Counter; print(Counter(t["description"] for t in tools).most_common(40))
